import io
import unittest

from docx import Document
from pptx import Presentation

from app.services.skill_exports import (
    render_document_docx,
    render_document_pdf,
    render_presentation_pptx,
)
from app.services.skill_generation import (
    _fallback_content,
    _normalize_document,
    _normalize_presentation,
)
from app.services.skill_registry import list_skills


class SkillsTests(unittest.TestCase):
    def test_catalogue_exposes_three_ready_skills(self):
        skills = list_skills()
        self.assertEqual(
            [item["id"] for item in skills],
            ["resume-builder", "document-writer", "presentation-builder"],
        )
        self.assertEqual(skills[0]["output_formats"], ["pdf", "docx"])
        self.assertEqual(skills[2]["output_formats"], ["pptx"])

    def test_document_normalizer_limits_and_cleans_sections(self):
        content = _normalize_document(
            {
                "title": "Executive brief",
                "sections": [
                    {"heading": "Outcome", "paragraphs": ["Clear result"], "bullets": ["One"]}
                ],
            },
            "",
            "executive-brief",
        )
        self.assertEqual(content["kind"], "document")
        self.assertEqual(content["sections"][0]["bullets"], ["One"])

    def test_presentation_normalizer_limits_slide_bullets(self):
        content = _normalize_presentation(
            {
                "title": "Plan",
                "slides": [{"title": "First 30 days", "bullets": [str(i) for i in range(12)]}],
            },
            "",
            "30-60-90",
        )
        self.assertEqual(content["kind"], "presentation")
        self.assertEqual(len(content["slides"][0]["bullets"]), 7)

    def test_document_exports_are_valid_files(self):
        content = _fallback_content(
            "document-writer", "executive-brief", "Test brief", "Grounded body"
        )
        docx_data = render_document_docx(content)
        pdf_data = render_document_pdf(content)
        parsed = Document(io.BytesIO(docx_data))
        self.assertIn("Test brief", "\n".join(p.text for p in parsed.paragraphs))
        self.assertTrue(pdf_data.startswith(b"%PDF"))

    def test_presentation_export_is_valid_widescreen_deck(self):
        content = _fallback_content(
            "presentation-builder", "interview-deck", "Interview plan", "Role context"
        )
        pptx_data = render_presentation_pptx(content)
        parsed = Presentation(io.BytesIO(pptx_data))
        self.assertEqual(len(parsed.slides), 4)
        self.assertGreater(parsed.slide_width, parsed.slide_height)


if __name__ == "__main__":
    unittest.main()
